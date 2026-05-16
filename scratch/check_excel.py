import openpyxl

path = r"c:\Work\project automation\coretax\MONITORING PK-PM RTSP 2026AAA.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

print(f"Sheet Name: {ws.title}")
print("-" * 50)

# Check Header
header_row = 1
for r in range(1, 15):
    row_vals = [cell.value for cell in ws[r]]
    if any(str(v).strip().upper() == "NOMOR FAKTUR" for v in row_vals if v):
        header_row = r
        print(f"Header Row detected at: {header_row}")
        print(f"Header values: {row_vals}")
        break

# Show some data rows (especially 'penjabaran' or columns that might be causing Qty=500)
# Let's find 'Harga' columns
harga_cols = []
for i, cell in enumerate(ws[header_row]):
    val = str(cell.value).upper() if cell.value else ""
    if "HARGA" in val:
        harga_cols.append((i+1, cell.value))

print(f"Harga Columns: {harga_cols}")

# Check rows starting from header_row + 1
for r in range(header_row + 1, header_row + 20):
    row_vals = [cell.value for cell in ws[r]]
    if any(row_vals):
        print(f"Row {r}: {row_vals}")

wb.close()
