import openpyxl

path = r"c:\Work\project automation\coretax\MONITORING PK-PM RTSP 2026AAA.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

header_row = 1
for r in range(1, 15):
    row_vals = [str(cell.value).strip().upper() if cell.value else "" for cell in ws[r]]
    if "NOMOR FAKTUR" in row_vals:
        header_row = r
        break

print(f"Header Row: {header_row}")
headers = [str(cell.value).strip() if cell.value else "" for cell in ws[header_row]]
for i, h in enumerate(headers):
    print(f"Col {i+1}: {h}")

wb.close()
