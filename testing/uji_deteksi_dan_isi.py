import openpyxl
import pandas as pd
import os

FILE_NAME = 'MONITORING PK-PM RTSP 2026AAA.xlsx'
SHEET_NAME = 'PM'

def run_test():
    print(f"[*] Membuka file: {FILE_NAME}")
    if not os.path.exists(FILE_NAME):
        print("[!] File tidak ditemukan!")
        return

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb[SHEET_NAME]
    
    # 1. Deteksi Header
    data_raw = list(ws.values)
    h_idx = 0
    for i, row in enumerate(data_raw):
        row_str = [str(v).upper() if v else "" for v in row]
        if any("NOMOR FAKTUR" in v for v in row_str):
            h_idx = i
            break
    
    headers = [str(c).strip() if c else f"COL_{i}" for i, c in enumerate(data_raw[h_idx])]
    print(f"[+] Header ditemukan di Baris: {h_idx + 1}")
    
    # 2. Cari Kolom Penting
    c_faktur = next((i+1 for i, c in enumerate(headers) if "NOMOR FAKTUR" in c.upper()), None)
    c_harga = next((i+1 for i, c in enumerate(headers) if "HARGA JUAL" in c.upper()), None)
    c_pj = next((i+1 for i, c in enumerate(headers) if "PENJABARAN" in c.upper()), None)
    c_bbm = next((i+1 for i, c in enumerate(headers) if "BBM" in c.upper() and "QTY" not in c.upper()), None)
    
    print(f"[i] Kolom Faktur: {c_faktur}, Kolom BBM: {c_bbm}, Kolom Penjabaran: {c_pj}")

    target_rows = [3, 5, 134]
    
    print("\n[*] TAHAP DETEKSI BARIS TARGET:")
    for r_idx in target_rows:
        faktur = ws.cell(row=r_idx, column=c_faktur).value
        pj_val = ws.cell(row=r_idx, column=c_pj).value
        print(f"--- Baris {r_idx} ---")
        print(f"    Faktur    : {faktur}")
        print(f"    Penjabaran: {pj_val}")
        
        # 3. COBA ISI ASAL (PROTOTYPE)
        print(f"    [!] Mencoba mengisi Baris {r_idx} dengan data asal...")
        
        # Isi Kolom BBM (sebagai contoh kategori)
        if c_bbm:
            ws.cell(row=r_idx, column=c_bbm).value = 888888
            ws.cell(row=r_idx, column=c_bbm).number_format = '#,##0'
        
        # Isi Kolom Penjabaran dengan RUMUS SUM
        if c_pj:
            # Kita asumsikan kategori ada di rentang K sampai AZ (contoh)
            # Tapi kita pakai kolom BBM saja untuk test
            col_letter = openpyxl.utils.get_column_letter(c_bbm)
            ws.cell(row=r_idx, column=c_pj).value = f"={col_letter}{r_idx}"
            ws.cell(row=r_idx, column=c_pj).number_format = '#,##0'
            print(f"    [OK] Rumus dipasang: ={col_letter}{r_idx}")

    # 4. Simpan
    try:
        wb.save(FILE_NAME)
        print("\n[√] BERHASIL! File sudah diupdate. Silakan cek Baris 3, 5, 134 di Excel Bapak.")
    except Exception as e:
        print(f"\n[!] GAGAL SIMPAN: {e} (Pastikan Excel sedang tidak dibuka!)")

if __name__ == "__main__":
    run_test()
